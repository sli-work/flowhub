"""附件解析器：数据结构 + 纯文本 / DOCX / XLSX / PPTX 的解析与证据定位。"""
from io import BytesIO

from flowhub_api.models.support import DocItem
from flowhub_api.services.attachment_parsers import parse_attachment

TXT = "需求: 备件库存看板\n数据来自 Q2 复盘。\n"


def _doc(name: str, doc_id: str = "dtest") -> DocItem:
    return DocItem(id=doc_id, name=name, project="p", version="v1", level="L2",
                   scan="已扫描", uploader="t", size="1KB", time="now")


async def test_plain_text_parsed_with_location():
    parsed = await parse_attachment(_doc("notes.txt"), TXT.encode("utf-8"))
    assert parsed.status == "indexed"
    assert parsed.parser == "plain"
    assert parsed.chunks and parsed.chunks[0].location.startswith("p1")
    assert "备件库存看板" in parsed.chunks[0].text


async def test_docx_parsed_with_paragraph_and_table():
    import docx

    document = docx.Document()
    document.add_heading("标题", level=1)
    document.add_paragraph("第一段正文")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "sku"
    table.cell(0, 1).text = "qty"
    table.cell(1, 0).text = "A-1"
    table.cell(1, 1).text = "42"
    buf = BytesIO()
    document.save(buf)
    parsed = await parse_attachment(_doc("doc.docx"), buf.getvalue())
    assert parsed.status == "indexed"
    assert parsed.parser == "docx"
    assert any(c.kind == "title" and "标题" in c.text for c in parsed.chunks)
    assert any(c.kind == "table" and "A-1" in c.text for c in parsed.chunks)


async def test_xlsx_parsed_with_sheet_location():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Q2"
    ws.append(["月份", "销售额"])
    ws.append(["4月", "100"])
    buf = BytesIO()
    wb.save(buf)
    parsed = await parse_attachment(_doc("book.xlsx"), buf.getvalue())
    assert parsed.status == "indexed"
    assert parsed.parser == "xlsx"
    assert any(c.location.startswith("Q2!") for c in parsed.chunks)
    assert any("销售额" in c.text for c in parsed.chunks)


async def test_pptx_parsed_with_slide_location():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "盘点报告"
    buf = BytesIO()
    prs.save(buf)
    parsed = await parse_attachment(_doc("deck.pptx"), buf.getvalue())
    assert parsed.status == "indexed"
    assert parsed.parser == "pptx"
    assert any(c.location.startswith("slide") for c in parsed.chunks)
    assert any("盘点报告" in c.text for c in parsed.chunks)


async def test_unsupported_extension_skipped():
    parsed = await parse_attachment(_doc("a.exe"), b"MZ....")
    assert parsed.status == "skipped"


async def test_pdf_text_page_parsed_with_location():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "备件库存看板", fontname="china-s")
    data = pdf.tobytes()
    pdf.close()
    parsed = await parse_attachment(_doc("doc.pdf"), data)
    assert parsed.status == "indexed"
    assert parsed.parser == "pdf"
    assert parsed.chunks and parsed.chunks[0].location.startswith("p")
    assert "备件库存看板" in parsed.chunks[0].text


async def test_pdf_no_text_without_ocr_needs_ocr():
    import fitz

    from flowhub_api.services.ocr import NoopOcrAdapter

    pdf = fitz.open()
    pdf.new_page()  # 空白页，无文本层
    data = pdf.tobytes()
    pdf.close()
    parsed = await parse_attachment(_doc("scan.pdf"), data, ocr=NoopOcrAdapter())
    assert parsed.status == "needs_ocr"
    assert any(c.kind == "ocr_marker" for c in parsed.chunks)


async def test_pdf_no_text_with_ocr_extracts_text():
    import fitz

    pdf = fitz.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()

    class FakeOcr:
        available = True

        async def extract_text(self, image_bytes, page_no):
            return "OCR 识别内容"

    parsed = await parse_attachment(_doc("scan.pdf"), data, ocr=FakeOcr())
    assert parsed.status == "indexed"
    assert any("OCR 识别内容" in c.text for c in parsed.chunks)


import zipfile


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, content in entries.items():
            zf.writestr(name, content)
    return buf.getvalue()


async def test_zip_indexes_whitelist_text_only():
    data = _zip_bytes({
        "readme.txt": "说明：Q2 备件盘点",
        "data/items.json": '{"sku": "A-1", "qty": 42}',
        "app.js": "var x = 1;",            # 脚本正文不进模型
        "bin.dat": b"\x00\x01\x02",         # 二进制不进模型
    })
    parsed = await parse_attachment(_doc("pkg.zip"), data)
    assert parsed.status == "indexed"
    assert parsed.parser == "zip"
    texts = "\n".join(c.text for c in parsed.chunks)
    assert "Q2 备件盘点" in texts
    assert "A-1" in texts
    assert "var x = 1" not in texts
    locations = {c.location for c in parsed.chunks}
    assert "readme.txt" in locations and "data/items.json" in locations


async def test_zip_nested_archive_skipped():
    inner = _zip_bytes({"inner.txt": "内部内容"})
    data = _zip_bytes({"outer.zip": inner, "top.txt": "顶层内容"})
    parsed = await parse_attachment(_doc("nested.zip"), data)
    assert parsed.status == "indexed"
    texts = "\n".join(c.text for c in parsed.chunks)
    assert "顶层内容" in texts
    assert "内部内容" not in texts


async def test_zip_axure_extracts_pages_and_title():
    data = _zip_bytes({
        "index.html": "<html><head><title>订单系统原型</title></head><body></body></html>",
        "data/document.js": 'var document = {"pages": [{"name": "首页", "id": "p1"}, {"name": "订单详情", "id": "p2"}]};',
        "data/document.css": "body { color: red; }",
        "js/script.js": "console.log('not for model');",
    })
    parsed = await parse_attachment(_doc("axure.zip"), data)
    assert parsed.status == "indexed"
    assert parsed.parser == "axure"
    texts = "\n".join(c.text for c in parsed.chunks)
    assert "订单系统原型" in texts
    assert "首页" in texts and "订单详情" in texts
    assert "console.log" not in texts


async def test_zip_path_traversal_rejected():
    data = _zip_bytes({"../evil.txt": "越权内容", "ok.txt": "正常"})
    parsed = await parse_attachment(_doc("bad.zip"), data)
    assert parsed.status == "failed"
    assert "路径" in parsed.error or "非法" in parsed.error


async def test_zip_compression_bomb_rejected():
    # 2000 个条目正好触顶，2001 个应拒绝
    entries = {f"f{i}.txt": b"x" * 100 for i in range(2001)}
    data = _zip_bytes(entries)
    parsed = await parse_attachment(_doc("bomb.zip"), data)
    assert parsed.status == "failed"
